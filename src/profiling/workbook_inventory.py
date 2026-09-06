from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


AUDIT_SIGNATURE = {
    "Clause ISO",
    "Libellé Terrain",
    "Constat de l'Auditeur",
}


DERIVED_REPORTING_SHEETS = {
    "DASHBOARD AUDITS",
    "CONSOLIDATION",
    "ANALYSE FINDINGS",
    "PLAN ACTION GLOBAL",
    "GRAPHIQUES",
}

DOCUMENTATION_SHEETS = {
    "ACCUEIL",
    "Revision History",
    "Legal Statement",
}

BUSINESS_RULES_SHEETS = {
    "REGLES DE GESTION",
}


def normalize(value: object) -> str:
    """
    Convertit une valeur Excel en chaîne nettoyée.
    """
    if value is None:
        return ""

    return str(value).strip()


def detect_header_row(
    worksheet,
    max_rows_to_scan: int = 25,
) -> tuple[int | None, list[str]]:
    """
    Recherche la ligne contenant les colonnes principales
    d'une checklist d'audit.

    Retourne :
        - numéro de la ligne d'en-tête ;
        - liste des colonnes détectées.
    """

    rows_to_scan = min(worksheet.max_row, max_rows_to_scan)

    for row_number in range(1, rows_to_scan + 1):

        values = [
            normalize(cell.value)
            for cell in worksheet[row_number]
        ]

        non_empty_values = {
            value
            for value in values
            if value
        }

        if AUDIT_SIGNATURE.issubset(non_empty_values):
            return row_number, values

    return None, []


def classify_sheet(
    sheet_name: str,
    header_row: int | None,
) -> str:
    """
    Classe une feuille selon son rôle dans le classeur.
    """

    if header_row is not None:
        return "audit_account"

    if sheet_name in BUSINESS_RULES_SHEETS:
        return "business_rules"

    if sheet_name in DERIVED_REPORTING_SHEETS:
        return "derived_reporting"

    if sheet_name in DOCUMENTATION_SHEETS:
        return "documentation"

    return "unknown_or_technical"


def profile_workbook(excel_path: Path) -> dict:
    """
    Construit l'inventaire structurel du classeur Excel.
    """

    workbook = load_workbook(
        excel_path,
        read_only=True,
        data_only=False,
    )

    inventory = {
        "source_file": excel_path.name,
        "sheet_count": len(workbook.sheetnames),
        "sheets": [],
    }

    for worksheet in workbook.worksheets:

        header_row, headers = detect_header_row(worksheet)

        sheet_type = classify_sheet(
            sheet_name=worksheet.title,
            header_row=header_row,
        )

        sheet_info = {
            "sheet_name": worksheet.title,
            "max_rows": worksheet.max_row,
            "max_columns": worksheet.max_column,
            "sheet_type": sheet_type,
            "audit_header_row": header_row,
            "headers": [
                header
                for header in headers
                if header
            ],
        }

        inventory["sheets"].append(sheet_info)

    workbook.close()

    return inventory


def save_inventory(
    inventory: dict,
    output_path: Path,
) -> None:
    """
    Sauvegarde le résultat du profiling au format JSON.
    """

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with output_path.open(
        "w",
        encoding="utf-8",
    ) as file:

        json.dump(
            inventory,
            file,
            ensure_ascii=False,
            indent=2,
        )


def parse_arguments():
    parser = argparse.ArgumentParser(
        description="Inventory and classify an Excel QHC workbook."
    )

    parser.add_argument(
        "--input",
        required=True,
        type=Path,
        help="Path to the source Excel workbook.",
    )

    parser.add_argument(
        "--output",
        default=Path("data/profiling/workbook_inventory.json"),
        type=Path,
        help="Path to the generated inventory JSON.",
    )

    return parser.parse_args()


def main():
    args = parse_arguments()

    if not args.input.exists():
        raise FileNotFoundError(
            f"Excel file not found: {args.input}"
        )

    inventory = profile_workbook(args.input)

    save_inventory(
        inventory=inventory,
        output_path=args.output,
    )

    print("Workbook profiling completed.")
    print(f"Sheets detected: {inventory['sheet_count']}")
    print(f"Inventory saved to: {args.output}")


if __name__ == "__main__":
    main()