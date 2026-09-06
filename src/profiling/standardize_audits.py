from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from openpyxl import load_workbook

from src.profiling.column_mapping import COLUMN_MAPPING
from src.profiling.workbook_inventory import detect_header_row


TECHNICAL_COLUMNS = [
    "source_file",
    "source_sheet",
]


def load_inventory(inventory_path: Path) -> dict:
    """
    Charge l'inventaire du classeur.
    """

    if not inventory_path.exists():
        raise FileNotFoundError(
            f"Inventory file not found: {inventory_path}"
        )

    with inventory_path.open(
        "r",
        encoding="utf-8",
    ) as file:
        return json.load(file)


def get_audit_sheet_names(inventory: dict) -> list[str]:
    """
    Retourne les noms des feuilles identifiées comme audits.
    """

    return [
        sheet["sheet_name"]
        for sheet in inventory["sheets"]
        if sheet["sheet_type"] == "audit_account"
    ]


def get_canonical_columns() -> list[str]:
    """
    Construit la liste unique des colonnes canoniques
    définies dans COLUMN_MAPPING.
    """

    return list(dict.fromkeys(COLUMN_MAPPING.values()))


def build_header_mapping(headers: list[str]) -> dict[int, str]:
    """
    Associe la position d'une colonne Excel
    à son nom canonique.

    Exemple :
        colonne Excel 12 : "Date de début"
        devient
        colonne canonique : "date_debut"
    """

    header_mapping = {}

    for column_index, source_header in enumerate(
        headers,
        start=1,
    ):
        if source_header in COLUMN_MAPPING:
            header_mapping[column_index] = COLUMN_MAPPING[
                source_header
            ]

    return header_mapping


def extract_standardized_rows(
    worksheet,
    source_file: str,
) -> list[dict]:
    """
    Extrait les lignes d'une feuille d'audit
    dans le schéma canonique.
    """

    header_row, headers = detect_header_row(worksheet)

    if header_row is None:
        raise ValueError(
            f"Audit header not found in sheet: {worksheet.title}"
        )

    header_mapping = build_header_mapping(headers)

    canonical_columns = get_canonical_columns()

    standardized_rows = []

    for excel_row_number in range(
        header_row + 1,
        worksheet.max_row + 1,
    ):
        row = {
            column: None
            for column in canonical_columns
        }

        row["source_file"] = source_file
        row["source_sheet"] = worksheet.title

        for column_index, canonical_name in header_mapping.items():

            value = worksheet.cell(
                row=excel_row_number,
                column=column_index,
            ).value

            row[canonical_name] = value

        # On conserve uniquement les vraies lignes
        # correspondant à la checklist d'audit.
        if not row.get("clause_iso") and not row.get(
            "libelle_terrain"
        ):
            continue

        standardized_rows.append(row)

    return standardized_rows


def standardize_workbook(
    excel_path: Path,
    inventory_path: Path,
) -> list[dict]:
    """
    Standardise toutes les feuilles d'audit du classeur.
    """

    inventory = load_inventory(inventory_path)

    audit_sheet_names = get_audit_sheet_names(inventory)

    workbook = load_workbook(
        excel_path,
        read_only=True,
        data_only=True,
    )

    standardized_rows = []

    for sheet_name in audit_sheet_names:

        worksheet = workbook[sheet_name]

        sheet_rows = extract_standardized_rows(
            worksheet=worksheet,
            source_file=excel_path.name,
        )

        standardized_rows.extend(sheet_rows)

        print(
            f"{sheet_name}: "
            f"{len(sheet_rows)} standardized rows"
        )

    workbook.close()

    return standardized_rows


def save_standardized_csv(
    rows: list[dict],
    output_path: Path,
) -> None:
    """
    Sauvegarde les données standardisées au format CSV.
    """

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    fieldnames = (
        TECHNICAL_COLUMNS
        + get_canonical_columns()
    )

    # Évite d'ajouter deux fois source_file/source_sheet
    fieldnames = list(dict.fromkeys(fieldnames))

    with output_path.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=fieldnames,
        )

        writer.writeheader()
        writer.writerows(rows)


def parse_arguments():
    parser = argparse.ArgumentParser(
        description="Standardize QHC audit sheets."
    )

    parser.add_argument(
        "--input",
        required=True,
        type=Path,
        help="Source QHC Excel workbook.",
    )

    parser.add_argument(
        "--inventory",
        default=Path(
            "data/profiling/workbook_inventory.json"
        ),
        type=Path,
    )

    parser.add_argument(
        "--output",
        default=Path(
            "data/standardized/standardized_audits.csv"
        ),
        type=Path,
    )

    return parser.parse_args()


def main():
    args = parse_arguments()

    if not args.input.exists():
        raise FileNotFoundError(
            f"Excel file not found: {args.input}"
        )

    rows = standardize_workbook(
        excel_path=args.input,
        inventory_path=args.inventory,
    )

    save_standardized_csv(
        rows=rows,
        output_path=args.output,
    )

    print("\nStandardization completed.")
    print(f"Total standardized rows: {len(rows)}")
    print(f"Output: {args.output}")


if __name__ == "__main__":
    main()